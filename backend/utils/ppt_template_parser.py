"""
PPT 模板解析器
用于分析 PPT 模板结构，识别可填充区域
"""
from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE_TYPE
from typing import Dict, List, Any, Optional
import os


class PPTTemplateParser:
    """PPT 模板解析器"""
    
    def __init__(self, template_path: str):
        """
        初始化模板解析器
        
        Args:
            template_path: 模板文件路径
        """
        if not os.path.exists(template_path):
            raise FileNotFoundError(f"模板文件不存在: {template_path}")
        
        self.template_path = template_path
        self.prs = Presentation(template_path)
        self.structure = self.analyze_structure()
    
    def analyze_structure(self) -> Dict[str, Any]:
        """
        分析模板结构
        
        Returns:
            模板结构信息
        """
        structure = {
            "total_slides": len(self.prs.slides),
            "slide_width": self.prs.slide_width,
            "slide_height": self.prs.slide_height,
            "slides": []
        }
        
        for idx, slide in enumerate(self.prs.slides):
            slide_info = self._analyze_slide(idx, slide)
            structure["slides"].append(slide_info)
        
        return structure
    
    def _analyze_slide(self, idx: int, slide) -> Dict[str, Any]:
        """分析单个幻灯片"""
        slide_info = {
            "index": idx,
            "layout_name": slide.slide_layout.name,
            "placeholders": [],
            "text_boxes": [],
            "images": [],
            "shapes": []
        }
        
        for shape in slide.shapes:
            # 识别占位符
            if shape.is_placeholder:
                slide_info["placeholders"].append({
                    "type": str(shape.placeholder_format.type),
                    "name": shape.name,
                    "text": shape.text if hasattr(shape, 'text') else None,
                    "left": shape.left,
                    "top": shape.top,
                    "width": shape.width,
                    "height": shape.height
                })
            
            # 识别文本框
            elif shape.has_text_frame:
                slide_info["text_boxes"].append({
                    "name": shape.name,
                    "text": shape.text,
                    "left": shape.left,
                    "top": shape.top,
                    "width": shape.width,
                    "height": shape.height
                })
            
            # 识别图片
            elif shape.shape_type == MSO_SHAPE_TYPE.PICTURE:
                slide_info["images"].append({
                    "name": shape.name,
                    "left": shape.left,
                    "top": shape.top,
                    "width": shape.width,
                    "height": shape.height
                })
            
            # 其他形状
            else:
                slide_info["shapes"].append({
                    "name": shape.name,
                    "type": str(shape.shape_type),
                    "left": shape.left,
                    "top": shape.top,
                    "width": shape.width,
                    "height": shape.height
                })
        
        return slide_info
    
    def get_fillable_areas(self, slide_index: int) -> Dict[str, Any]:
        """
        获取可填充区域
        
        Args:
            slide_index: 幻灯片索引
            
        Returns:
            可填充区域信息
        """
        if slide_index >= len(self.structure["slides"]):
            return {"title": None, "content": [], "images": []}
        
        slide_info = self.structure["slides"][slide_index]
        
        return {
            "title": self._find_title_area(slide_info),
            "content": self._find_content_areas(slide_info),
            "images": slide_info["images"]
        }
    
    def _find_title_area(self, slide_info: Dict) -> Optional[Dict]:
        """查找标题区域"""
        # 优先查找标题占位符
        for ph in slide_info["placeholders"]:
            if "TITLE" in ph["type"] or "title" in ph["name"].lower():
                return ph
        
        # 查找最上方的文本框（通常是标题）
        text_boxes = slide_info["text_boxes"]
        if text_boxes:
            # 按 top 位置排序，取最上方的
            sorted_boxes = sorted(text_boxes, key=lambda x: x["top"])
            return sorted_boxes[0]
        
        return None
    
    def _find_content_areas(self, slide_info: Dict) -> List[Dict]:
        """查找内容区域"""
        content_areas = []
        
        # 查找内容占位符
        for ph in slide_info["placeholders"]:
            ph_type = ph["type"]
            ph_name = ph["name"].lower()
            
            if ("BODY" in ph_type or "OBJECT" in ph_type or 
                "content" in ph_name or "body" in ph_name or "text" in ph_name):
                content_areas.append(ph)
        
        # 如果没有找到占位符，查找其他文本框（排除标题）
        if not content_areas:
            title = self._find_title_area(slide_info)
            for tb in slide_info["text_boxes"]:
                if title and tb["name"] != title["name"]:
                    content_areas.append(tb)
        
        return content_areas
    
    def get_summary(self) -> str:
        """获取模板摘要信息"""
        summary = f"模板信息:\n"
        summary += f"- 总页数: {self.structure['total_slides']}\n"
        summary += f"- 尺寸: {self.structure['slide_width']} x {self.structure['slide_height']}\n\n"
        
        for idx, slide in enumerate(self.structure["slides"]):
            summary += f"第 {idx + 1} 页:\n"
            summary += f"  布局: {slide['layout_name']}\n"
            summary += f"  占位符: {len(slide['placeholders'])} 个\n"
            summary += f"  文本框: {len(slide['text_boxes'])} 个\n"
            summary += f"  图片: {len(slide['images'])} 个\n\n"
        
        return summary


if __name__ == "__main__":
    # 测试代码
    import sys
    
    if len(sys.argv) > 1:
        template_path = sys.argv[1]
        parser = PPTTemplateParser(template_path)
        print(parser.get_summary())
