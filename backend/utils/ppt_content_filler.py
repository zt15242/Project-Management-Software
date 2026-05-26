"""
PPT 内容填充器
用于将 AI 生成的内容填充到 PPT 模板中
"""
from pptx import Presentation
from pptx.util import Pt
from pptx.enum.shapes import MSO_SHAPE_TYPE
from typing import Dict, List, Any, Optional
import io
import os
import time
from utils.ppt_template_parser import PPTTemplateParser


class PPTContentFiller:
    """PPT 内容填充器"""
    
    def __init__(self, template_path: str):
        """
        初始化内容填充器
        
        Args:
            template_path: 模板文件路径
        """
        self.template_path = template_path
        self.parser = PPTTemplateParser(template_path)
    
    async def fill_content(
        self, 
        ai_content: Dict[str, Any],
        use_ai_images: bool = True,
        ai_config: Optional[Dict] = None
    ) -> str:
        """
        填充内容到模板
        
        Args:
            ai_content: AI 生成的内容
                {
                    "title": "PPT标题",
                    "slides": [
                        {
                            "title": "页面标题",
                            "content": ["要点1", "要点2"],
                            "image_description": "图片描述"
                        }
                    ]
                }
            use_ai_images: 是否使用 AI 生成图片
            ai_config: AI 配置（用于图片生成）
            
        Returns:
            生成的 PPT 文件路径
        """
        print(f"[Template Filler] 开始填充内容到模板: {self.template_path}")
        
        # 加载模板
        prs = Presentation(self.template_path)
        
        # 填充每一页
        slides_to_fill = min(len(ai_content["slides"]), len(prs.slides))
        print(f"[Template Filler] 将填充 {slides_to_fill} 页内容")
        
        for idx in range(slides_to_fill):
            slide_content = ai_content["slides"][idx]
            slide = prs.slides[idx]
            fillable = self.parser.get_fillable_areas(idx)
            
            print(f"[Template Filler] 填充第 {idx + 1} 页")
            
            # 填充标题
            if fillable["title"] and slide_content.get("title"):
                self._fill_title(slide, fillable["title"], slide_content["title"])
            
            # 填充内容
            if fillable["content"] and slide_content.get("content"):
                self._fill_content(slide, fillable["content"], slide_content["content"])
            
            # 替换图片
            if use_ai_images and fillable["images"] and ai_config:
                await self._replace_images(
                    slide, 
                    fillable["images"], 
                    slide_content.get("image_description", ""),
                    ai_config
                )
        
        # 保存
        output_dir = os.path.join(os.path.dirname(self.template_path), "..", "generated")
        os.makedirs(output_dir, exist_ok=True)
        
        output_filename = f"ppt_{int(time.time())}.pptx"
        output_path = os.path.join(output_dir, output_filename)
        
        prs.save(output_path)
        print(f"[Template Filler] PPT 已保存: {output_path}")
        
        return output_path
    
    def _fill_title(self, slide, title_area: Dict, text: str):
        """填充标题"""
        print(f"[Template Filler] 填充标题: {text[:30]}...")
        
        for shape in slide.shapes:
            if shape.name == title_area["name"]:
                if shape.has_text_frame:
                    # 保存原有格式
                    original_font_size = None
                    original_font_bold = None
                    
                    if shape.text_frame.paragraphs:
                        first_para = shape.text_frame.paragraphs[0]
                        if first_para.runs:
                            first_run = first_para.runs[0]
                            original_font_size = first_run.font.size
                            original_font_bold = first_run.font.bold
                    
                    # 设置新文本
                    shape.text = text
                    
                    # 恢复格式
                    if shape.text_frame.paragraphs:
                        for paragraph in shape.text_frame.paragraphs:
                            for run in paragraph.runs:
                                if original_font_size:
                                    run.font.size = original_font_size
                                if original_font_bold is not None:
                                    run.font.bold = original_font_bold
                break
    
    def _fill_content(self, slide, content_areas: List[Dict], content_list: List[str]):
        """填充内容"""
        if not content_areas or not content_list:
            return
        
        print(f"[Template Filler] 填充 {len(content_list)} 条内容")
        
        # 使用第一个内容区域
        content_area = content_areas[0]
        
        for shape in slide.shapes:
            if shape.name == content_area["name"]:
                if shape.has_text_frame:
                    tf = shape.text_frame
                    
                    # 保存原有格式
                    original_font_size = None
                    if tf.paragraphs and tf.paragraphs[0].runs:
                        original_font_size = tf.paragraphs[0].runs[0].font.size
                    
                    # 清空原有内容
                    tf.clear()
                    
                    # 添加新内容
                    for i, item in enumerate(content_list):
                        p = tf.add_paragraph()
                        p.text = f"• {item}"
                        p.level = 0
                        
                        # 恢复字体大小
                        if original_font_size and p.runs:
                            for run in p.runs:
                                run.font.size = original_font_size
                break
    
    async def _replace_images(
        self, 
        slide, 
        image_areas: List[Dict], 
        description: str,
        ai_config: Dict
    ):
        """替换图片为 AI 生成的图片"""
        if not description:
            print("[Template Filler] 没有图片描述，跳过图片生成")
            return
        
        print(f"[Template Filler] 准备生成图片: {description[:50]}...")
        
        from utils.gemini_image import get_image_generator
        
        # 生成图片
        image_gen = get_image_generator()
        
        # 为每个图片区域生成图片
        for img_area in image_areas:
            try:
                print(f"[Template Filler] 为图片区域 '{img_area['name']}' 生成图片")
                
                # 生成图片
                image_bytes = await image_gen.generate_ppt_slide_image(
                    slide_data={
                        "title": "",
                        "content": [],
                        "image_description": description,
                        "layout": "full"
                    },
                    theme_color="#1e3a8a",
                    style="professional",
                    ai_config=ai_config
                )
                
                if image_bytes:
                    # 替换图片
                    self._replace_image_in_slide(
                        slide, 
                        img_area, 
                        io.BytesIO(image_bytes)
                    )
                    print(f"[Template Filler] 图片替换成功")
                else:
                    print(f"[Template Filler] 图片生成失败，保留原图")
                    
            except Exception as e:
                print(f"[Template Filler] 图片替换失败: {str(e)}")
                import traceback
                traceback.print_exc()
    
    def _replace_image_in_slide(self, slide, img_area: Dict, image_stream):
        """在幻灯片中替换图片"""
        # 找到原图片
        shape_to_replace = None
        for shape in slide.shapes:
            if shape.name == img_area["name"]:
                shape_to_replace = shape
                break
        
        if not shape_to_replace:
            print(f"[Template Filler] 未找到图片: {img_area['name']}")
            return
        
        # 记录位置和大小
        left = shape_to_replace.left
        top = shape_to_replace.top
        width = shape_to_replace.width
        height = shape_to_replace.height
        
        # 删除原图片
        sp = shape_to_replace.element
        sp.getparent().remove(sp)
        
        # 添加新图片
        try:
            slide.shapes.add_picture(
                image_stream,
                left, top,
                width=width,
                height=height
            )
            print(f"[Template Filler] 图片已替换")
        except Exception as e:
            print(f"[Template Filler] 添加图片失败: {str(e)}")
