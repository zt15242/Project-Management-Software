from fastapi import APIRouter, Depends, HTTPException, status, UploadFile, File, Form
from fastapi.responses import FileResponse
from typing import List, Optional, Any
from database import get_database
from models import (
    KnowledgeCreate, KnowledgeUpdate, KnowledgeResponse, 
    KnowledgeFileType, KnowledgeStatus, KnowledgeSearchRequest,
    UserResponse, UserRole
)
from auth import get_current_active_user
from bson import ObjectId
from datetime import datetime, timedelta
from utils import get_beijing_time
from config import settings
import os
import aiofiles
import uuid
import mimetypes
import secrets
from utils.oss_utils import upload_file_to_oss, delete_file_from_oss

router = APIRouter(prefix="/api/knowledge", tags=["知识库"])

# 临时访问令牌存储（生产环境建议使用Redis）
_temp_tokens = {}


# 文件类型映射
FILE_TYPE_MAPPING = {
    '.md': KnowledgeFileType.MARKDOWN,
    '.markdown': KnowledgeFileType.MARKDOWN,
    '.doc': KnowledgeFileType.WORD,
    '.docx': KnowledgeFileType.WORD,
    '.xls': KnowledgeFileType.EXCEL,
    '.xlsx': KnowledgeFileType.EXCEL,
    '.pdf': KnowledgeFileType.PDF,
    '.txt': KnowledgeFileType.TEXT,
    '.png': KnowledgeFileType.IMAGE,
    '.jpg': KnowledgeFileType.IMAGE,
    '.jpeg': KnowledgeFileType.IMAGE,
    '.gif': KnowledgeFileType.IMAGE,
    '.ppt': KnowledgeFileType.PPT,
    '.pptx': KnowledgeFileType.PPT,
}


def get_file_type(filename: str) -> KnowledgeFileType:
    """根据文件名获取文件类型"""
    ext = os.path.splitext(filename)[1].lower()
    return FILE_TYPE_MAPPING.get(ext, KnowledgeFileType.OTHER)


async def extract_text_content(file_input: Any, file_type: KnowledgeFileType) -> Optional[str]:
    """
    提取文件的文本内容
    使用 asyncio.to_thread 将同步的、耗时的解析逻辑放到线程池中，避免阻塞主线程
    """
    import asyncio
    import io
    
    is_path = isinstance(file_input, str)
    
    # 不同的文件类型解析逻辑
    def _sync_parse():
        f_obj = None
        try:
            # 1. Markdown 和 Text
            if file_type == KnowledgeFileType.MARKDOWN or file_type == KnowledgeFileType.TEXT:
                if is_path:
                    with open(file_input, 'r', encoding='utf-8', errors='ignore') as f:
                        return f.read()
                else:
                    file_input.seek(0)
                    return file_input.read().decode('utf-8', errors='ignore')
            
            # 2. PDF
            elif file_type == KnowledgeFileType.PDF:
                import PyPDF2
                f_obj = open(file_input, 'rb') if is_path else file_input
                if not is_path: f_obj.seek(0)
                reader = PyPDF2.PdfReader(f_obj)
                text = ""
                for page in reader.pages:
                    text += page.extract_text() + "\n"
                return text
                
            # 3. Word
            elif file_type == KnowledgeFileType.WORD:
                import docx
                f_obj = open(file_input, 'rb') if is_path else file_input
                if not is_path: f_obj.seek(0)
                doc = docx.Document(f_obj)
                text_parts = []
                for para in doc.paragraphs:
                    if para.text.strip():
                        text_parts.append(para.text)
                for table in doc.tables:
                    for row in table.rows:
                        row_text = '\t'.join([cell.text for cell in row.cells])
                        if row_text.strip():
                            text_parts.append(row_text)
                return "\n".join(text_parts)
                
            # 4. Excel
            elif file_type == KnowledgeFileType.EXCEL:
                import openpyxl
                f_obj = open(file_input, 'rb') if is_path else file_input
                if not is_path: f_obj.seek(0)
                wb = openpyxl.load_workbook(f_obj, data_only=True)
                text = ""
                for sheet in wb.worksheets:
                    text += f"\n=== {sheet.title} ===\n"
                    for row in sheet.iter_rows(values_only=True):
                        text += "\t".join([str(cell) if cell is not None else "" for cell in row]) + "\n"
                return text
                
            return None
        except Exception as e:
            print(f"解析文件出错: {e}")
            return None
        finally:
            if is_path and f_obj: f_obj.close()

    # 将同步解析任务提交到线程池执行
    return await asyncio.to_thread(_sync_parse)


@router.post("/", response_model=KnowledgeResponse)
async def create_knowledge(
    title: str = Form(...),
    project_id: str = Form(...),
    category: Optional[str] = Form(None),
    description: Optional[str] = Form(None),
    tags: Optional[str] = Form(None),  # 逗号分隔的标签
    file: UploadFile = File(...),
    current_user: UserResponse = Depends(get_current_active_user)
):
    """创建知识库文档（上传文件）"""
    db = get_database()
    
    # 验证项目是否存在
    try:
        project = await db.projects.find_one({"_id": ObjectId(project_id)})
    except:
        raise HTTPException(status_code=400, detail="无效的项目ID")
    
    if not project:
        raise HTTPException(status_code=404, detail="项目不存在")
    
    # 检查权限
    if current_user.role != UserRole.ADMIN and current_user.id not in project["team_members"]:
        raise HTTPException(status_code=403, detail="没有权限")
    
    # 创建知识库目录
    knowledge_dir = os.path.join(settings.UPLOAD_DIR, "knowledge", project_id)
    os.makedirs(knowledge_dir, exist_ok=True)
    
    # 生成唯一文件名
    file_ext = os.path.splitext(file.filename)[1]
    unique_filename = f"{uuid.uuid4()}{file_ext}"
    file_path = os.path.join(knowledge_dir, unique_filename)
    
    # 保存文件
    async with aiofiles.open(file_path, 'wb') as f:
        content = await file.read()
        await f.write(content)
        file_size = len(content)
    
    
    # 获取文件类型
    file_type = get_file_type(file.filename)
    
    # 提取文本内容（在上传OSS之前，确保本地文件存在）
    text_content = await extract_text_content(file_path, file_type)
    
    # 尝试上传到OSS
    is_oss = False
    oss_url = None
    try:
        object_name = f"knowledge/{project_id}/{unique_filename}"
        oss_url = await upload_file_to_oss(file_path, object_name)
        
        # 检查是否上传到了OSS（如果是OSS，URL将不包含BASE_URL当中的uploads部分，或者我们显式判断）
        # upload_file_to_oss 如果未配置OSS会返回本地URL
        
        # 简单判断: 如果返回的不是本地路径格式(based on settings.BASE_URL)，则认为是OSS
        base_url = getattr(settings, 'BASE_URL', 'http://localhost:8000')
        is_oss = not oss_url.startswith(f"{base_url}/uploads")
        
        if is_oss:
            # 上传成功，删除本地文件
            try:
                os.remove(file_path)
                print(f"本地文件已删除: {file_path}")
            except Exception as e:
                print(f"删除本地临时文件失败: {e}")
            
    except Exception as e:
        print(f"上传OSS失败: {e}")
        # 如果失败，保留本地文件作为后备
    
    # 解析标签
    tag_list = []
    if tags:
        tag_list = [tag.strip() for tag in tags.split(",") if tag.strip()]
    
    # 创建知识库文档
    knowledge_doc = {
        "title": title,
        "project_id": project_id,
        "category": category,
        "description": description,
        "file_type": file_type,
        "file_name": file.filename,
        "file_path": oss_url if is_oss else file_path,  # 使用OSS URL或本地路径
        "file_size": file_size,
        "content": text_content,
        "tags": tag_list,
        "status": KnowledgeStatus.PUBLISHED,
        "views": 0,
        "created_by": current_user.id,
        "created_at": get_beijing_time(),
        "updated_at": get_beijing_time()
    }
    
    if is_oss:
        knowledge_doc["storage_type"] = "oss"
        knowledge_doc["oss_key"] = object_name
    else:
        knowledge_doc["storage_type"] = "local"
    
    result = await db.knowledge.insert_one(knowledge_doc)
    created_doc = await db.knowledge.find_one({"_id": result.inserted_id})
    
    return KnowledgeResponse(
        id=str(created_doc["_id"]),
        title=created_doc["title"],
        project_id=created_doc["project_id"],
        category=created_doc.get("category"),
        description=created_doc.get("description"),
        file_type=created_doc["file_type"],
        file_name=created_doc["file_name"],
        file_path=created_doc["file_path"],
        file_size=created_doc["file_size"],
        content=created_doc.get("content"),
        tags=created_doc.get("tags", []),
        status=created_doc["status"],
        views=created_doc.get("views", 0),
        created_by=created_doc["created_by"],
        created_at=created_doc["created_at"],
        updated_at=created_doc["updated_at"]
    )


@router.get("/", response_model=List[KnowledgeResponse])
async def get_knowledge_list(
    project_id: Optional[str] = None,
    category: Optional[str] = None,
    keyword: Optional[str] = None,
    file_type: Optional[KnowledgeFileType] = None,
    status: Optional[KnowledgeStatus] = None,
    current_user: UserResponse = Depends(get_current_active_user)
):
    """获取知识库文档列表（支持搜索和过滤）"""
    db = get_database()
    
    # 构建查询条件
    query = {}
    
    if project_id:
        query["project_id"] = project_id
    
    if category:
        query["category"] = category
    
    if file_type:
        query["file_type"] = file_type
    
    if status:
        query["status"] = status
    
    # 关键词搜索（标题、描述、内容）
    if keyword:
        query["$or"] = [
            {"title": {"$regex": keyword, "$options": "i"}},
            {"description": {"$regex": keyword, "$options": "i"}},
            {"content": {"$regex": keyword, "$options": "i"}}
        ]
    
    # 获取文档列表
    docs = await db.knowledge.find(query).sort("created_at", -1).to_list(length=None)
    
    return [
        KnowledgeResponse(
            id=str(doc["_id"]),
            title=doc["title"],
            project_id=doc["project_id"],
            category=doc.get("category"),
            description=doc.get("description"),
            file_type=doc["file_type"],
            file_name=doc["file_name"],
            file_path=doc["file_path"],
            file_size=doc["file_size"],
            content=doc.get("content"),
            tags=doc.get("tags", []),
            status=doc["status"],
            views=doc.get("views", 0),
            created_by=doc["created_by"],
            created_at=doc.get("created_at") or get_beijing_time(),
            updated_at=doc.get("updated_at") or get_beijing_time()
        )
        for doc in docs
    ]


@router.get("/{knowledge_id}", response_model=KnowledgeResponse)
async def get_knowledge(
    knowledge_id: str,
    current_user: UserResponse = Depends(get_current_active_user)
):
    """获取知识库文档详情"""
    db = get_database()
    
    try:
        doc = await db.knowledge.find_one({"_id": ObjectId(knowledge_id)})
    except:
        raise HTTPException(status_code=400, detail="无效的文档ID")
    
    if not doc:
        raise HTTPException(status_code=404, detail="文档不存在")
    
    # 增加浏览次数
    await db.knowledge.update_one(
        {"_id": ObjectId(knowledge_id)},
        {"$inc": {"views": 1}}
    )
    
    return KnowledgeResponse(
        id=str(doc["_id"]),
        title=doc["title"],
        project_id=doc["project_id"],
        category=doc.get("category"),
        description=doc.get("description"),
        file_type=doc["file_type"],
        file_name=doc["file_name"],
        file_path=doc["file_path"],
        file_size=doc["file_size"],
        content=doc.get("content"),
        tags=doc.get("tags", []),
        status=doc["status"],
        views=doc.get("views", 0) + 1,
        created_by=doc["created_by"],
        created_at=doc["created_at"],
        updated_at=doc["updated_at"]
    )


@router.put("/{knowledge_id}", response_model=KnowledgeResponse)
async def update_knowledge(
    knowledge_id: str,
    knowledge_update: KnowledgeUpdate,
    current_user: UserResponse = Depends(get_current_active_user)
):
    """更新知识库文档信息"""
    db = get_database()
    
    try:
        doc = await db.knowledge.find_one({"_id": ObjectId(knowledge_id)})
    except:
        raise HTTPException(status_code=400, detail="无效的文档ID")
    
    if not doc:
        raise HTTPException(status_code=404, detail="文档不存在")
    
    # 检查权限
    if current_user.role != UserRole.ADMIN and current_user.id != doc["created_by"]:
        raise HTTPException(status_code=403, detail="没有权限")
    
    # 构建更新数据
    update_data = {k: v for k, v in knowledge_update.dict(exclude_unset=True).items() if v is not None}
    
    if not update_data:
        raise HTTPException(status_code=400, detail="没有需要更新的数据")
    
    update_data["updated_at"] = get_beijing_time()
    
    # 更新文档
    await db.knowledge.update_one(
        {"_id": ObjectId(knowledge_id)},
        {"$set": update_data}
    )
    
    updated_doc = await db.knowledge.find_one({"_id": ObjectId(knowledge_id)})
    
    return KnowledgeResponse(
        id=str(updated_doc["_id"]),
        title=updated_doc["title"],
        project_id=updated_doc["project_id"],
        category=updated_doc.get("category"),
        description=updated_doc.get("description"),
        file_type=updated_doc["file_type"],
        file_name=updated_doc["file_name"],
        file_path=updated_doc["file_path"],
        file_size=updated_doc["file_size"],
        content=updated_doc.get("content"),
        tags=updated_doc.get("tags", []),
        status=updated_doc["status"],
        views=updated_doc.get("views", 0),
        created_by=updated_doc["created_by"],
        created_at=updated_doc["created_at"],
        updated_at=updated_doc["updated_at"]
    )


@router.delete("/{knowledge_id}", status_code=status.HTTP_204_NO_CONTENT)
async def delete_knowledge(
    knowledge_id: str,
    current_user: UserResponse = Depends(get_current_active_user)
):
    """删除知识库文档"""
    db = get_database()
    
    try:
        doc = await db.knowledge.find_one({"_id": ObjectId(knowledge_id)})
    except:
        raise HTTPException(status_code=400, detail="无效的文档ID")
    
    if not doc:
        raise HTTPException(status_code=404, detail="文档不存在")
    
    # 检查权限
    if current_user.role != UserRole.ADMIN and current_user.id != doc["created_by"]:
        raise HTTPException(status_code=403, detail="没有权限")
    
    # 删除文件
    # 删除文件
    try:
        # 检查是否是OSS文件
        if doc.get("storage_type") == "oss" or doc.get("file_path", "").startswith("http"):
            # 尝试从OSS删除
            oss_key = doc.get("oss_key")
            if not oss_key:
                # 尝试从URL中提取（如果不完美也没关系，主要是尝试清理）
                # 简单处理：如果这就是一个OSS URL，我们尝试用delete_file_from_oss
                # 但我们需要key。如果数据库没存oss_key，可能无法准确删除。
                pass
            
            if oss_key:
                await delete_file_from_oss(oss_key)
        
        # 尝试删除本地文件（兼容旧数据或本地存储）
        if os.path.exists(doc["file_path"]):
            os.remove(doc["file_path"])
    except Exception as e:
        print(f"删除文件失败: {e}")
    
    # 删除数据库记录
    await db.knowledge.delete_one({"_id": ObjectId(knowledge_id)})
    
    return None


@router.get("/categories/list")
async def get_categories(
    project_id: Optional[str] = None,
    current_user: UserResponse = Depends(get_current_active_user)
):
    """获取所有分类列表"""
    db = get_database()
    
    query = {}
    if project_id:
        query["project_id"] = project_id
    
    # 获取所有不同的分类
    categories = await db.knowledge.distinct("category", query)
    
    # 过滤掉None值
    categories = [cat for cat in categories if cat]
    
    return {"categories": categories}


@router.get("/tags/list")
async def get_tags(
    project_id: Optional[str] = None,
    current_user: UserResponse = Depends(get_current_active_user)
):
    """获取所有标签列表"""
    db = get_database()
    
    query = {}
    if project_id:
        query["project_id"] = project_id
    
    # 获取所有文档的标签
    docs = await db.knowledge.find(query, {"tags": 1}).to_list(length=None)
    
    # 合并所有标签并去重
    all_tags = set()
    for doc in docs:
        tags = doc.get("tags", [])
        all_tags.update(tags)
    
    return {"tags": sorted(list(all_tags))}


@router.get("/statistics/{project_id}")
async def get_knowledge_statistics(
    project_id: str,
    current_user: UserResponse = Depends(get_current_active_user)
):
    """获取知识库统计信息"""
    db = get_database()
    
    # 总文档数
    total_docs = await db.knowledge.count_documents({"project_id": project_id})
    
    # 按文件类型统计
    file_type_stats = {}
    for file_type in KnowledgeFileType:
        count = await db.knowledge.count_documents({
            "project_id": project_id,
            "file_type": file_type
        })
        if count > 0:
            file_type_stats[file_type] = count
    
    # 按状态统计
    status_stats = {}
    for status_val in KnowledgeStatus:
        count = await db.knowledge.count_documents({
            "project_id": project_id,
            "status": status_val
        })
        if count > 0:
            status_stats[status_val] = count
    
    # 总浏览次数
    pipeline = [
        {"$match": {"project_id": project_id}},
        {"$group": {"_id": None, "total_views": {"$sum": "$views"}}}
    ]
    result = await db.knowledge.aggregate(pipeline).to_list(length=1)
    total_views = result[0]["total_views"] if result else 0
    
    return {
        "total_docs": total_docs,
        "file_type_stats": file_type_stats,
        "status_stats": status_stats,
        "total_views": total_views
    }


@router.get("/{knowledge_id}/public-url")
async def generate_public_url(
    knowledge_id: str,
    expires_minutes: int = 10,  # 默认10分钟有效期
    current_user: UserResponse = Depends(get_current_active_user)
):
    """
    生成文档的临时公开访问URL
    用于Office Online等外部服务访问
    """
    db = get_database()
    
    try:
        doc = await db.knowledge.find_one({"_id": ObjectId(knowledge_id)})
    except:
        raise HTTPException(status_code=400, detail="无效的文档ID")
    
    if not doc:
        raise HTTPException(status_code=404, detail="文档不存在")
    
    # 生成临时令牌（使用secrets生成安全的随机字符串）
    temp_token = secrets.token_urlsafe(32)
    
    # 计算过期时间
    expires_at = datetime.now() + timedelta(minutes=expires_minutes)
    
    # 存储令牌信息
    _temp_tokens[temp_token] = {
        "knowledge_id": knowledge_id,
        "file_path": doc["file_path"],
        "file_name": doc["file_name"],
        "expires_at": expires_at
    }
    
    # 构建公开URL
    # 注意：这里需要使用服务器的公网地址
    # 如果settings中没有BASE_URL，使用默认值
    base_url = getattr(settings, 'BASE_URL', 'http://localhost:8000')
    
    # 如果是OSS文件，直接返回文件URL
    if doc.get("storage_type") == "oss" or (doc.get("file_path", "").startswith("http") and not doc.get("file_path", "").startswith(f"{base_url}/uploads")):
        public_url = doc["file_path"]
    else:
        public_url = f"{base_url}/api/knowledge/public/{temp_token}"
    
    return {
        "public_url": public_url,
        "temp_token": temp_token,
        "expires_at": expires_at.isoformat(),
        "expires_in_seconds": expires_minutes * 60
    }


@router.get("/public/{temp_token}")
async def get_public_file(temp_token: str):
    """
    通过临时令牌访问文件（无需认证）
    用于Office Online等外部服务
    """
    # 检查令牌是否存在
    if temp_token not in _temp_tokens:
        raise HTTPException(status_code=404, detail="无效的访问令牌")
    
    token_info = _temp_tokens[temp_token]
    
    # 检查是否过期
    if datetime.now() > token_info["expires_at"]:
        # 清理过期令牌
        del _temp_tokens[temp_token]
        raise HTTPException(status_code=410, detail="访问令牌已过期")
    
    file_path = token_info["file_path"]
    file_name = token_info["file_name"]
    
    # 检查文件是否存在
    if not os.path.exists(file_path):
        raise HTTPException(status_code=404, detail="文件不存在")
    
    # 返回文件
    # 设置正确的Content-Type以便Office Online识别
    media_type = mimetypes.guess_type(file_name)[0] or "application/octet-stream"
    
    return FileResponse(
        path=file_path,
        filename=file_name,
        media_type=media_type,
        headers={
            "Access-Control-Allow-Origin": "*",  # 允许Office Online跨域访问
            "Cache-Control": "no-cache"
        }
    )
