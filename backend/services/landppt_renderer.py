"""
LandPPT 模板渲染引擎
将 HTML 模板渲染为 PPTX 幻灯片
"""
import asyncio
import logging
from typing import Dict, Any, Optional, List
from pathlib import Path
from jinja2 import Template
from playwright.async_api import async_playwright
import io
from PIL import Image
from pptx import Presentation
from pptx.util import Inches
from datetime import datetime

logger = logging.getLogger(__name__)


class LandPPTTemplateRenderer:
    """LandPPT 模板渲染器"""
    
    def __init__(self, template_html: str):
        """
        初始化渲染器
        
        Args:
            template_html: HTML 模板字符串
        """
        self.template_html = template_html
        self.jinja_template = Template(template_html)
    
    def render_html(self, slide_data: Dict[str, Any]) -> str:
        """
        渲染 HTML
        
        Args:
            slide_data: 幻灯片数据，包含：
                - page_title: 页面标题
                - main_heading: 主标题
                - page_content: 页面内容（HTML）
                - current_page_number: 当前页码
                - total_page_count: 总页数
        
        Returns:
            渲染后的 HTML 字符串
        """
        try:
            rendered_html = self.jinja_template.render(**slide_data)
            return rendered_html
        except Exception as e:
            logger.error(f"HTML 渲染失败: {e}")
            raise
    
    async def html_to_image(self, html_content: str, output_path: Optional[str] = None) -> bytes:
        """
        将 HTML 转换为图片 (带同步回退机制)
        
        Args:
            html_content: HTML 内容
            output_path: 输出路径（可选）
        
        Returns:
            图片字节数据
        """
        try:
            # 尝试使用异步 Playwright
            async with async_playwright() as p:
                browser = await p.chromium.launch(headless=True)
                page = await browser.new_page(viewport={'width': 1280, 'height': 720})
                
                try:
                    await page.set_content(html_content, wait_until='networkidle')
                    # await asyncio.sleep(0.5) # 减少等待时间
                    
                    screenshot_bytes = await page.screenshot(type='png', full_page=False)
                    
                    if output_path:
                        with open(output_path, 'wb') as f:
                            f.write(screenshot_bytes)
                    
                    return screenshot_bytes
                finally:
                    await browser.close()
                    
        except NotImplementedError:
            # Windows asyncio 问题，回退到同步 Playwright (在线程中运行)
            logger.warning("捕获到 NotImplementedError，切换到同步 Playwright 模式")
            return await asyncio.to_thread(self._html_to_image_sync, html_content, output_path)
            
        except Exception as e:
            logger.error(f"异步渲染失败: {e}")
            # 其他错误也尝试同步模式
            if "EventLoop" in str(e) or "win32" in str(e):
                 return await asyncio.to_thread(self._html_to_image_sync, html_content, output_path)
            raise

    def _html_to_image_sync(self, html_content: str, output_path: Optional[str] = None) -> bytes:
        """同步版本的 HTML 转图片"""
        from playwright.sync_api import sync_playwright
        
        with sync_playwright() as p:
            browser = p.chromium.launch(headless=True)
            page = browser.new_page(viewport={'width': 1280, 'height': 720})
            
            try:
                page.set_content(html_content, wait_until='networkidle')
                # page.wait_for_timeout(500)
                
                screenshot_bytes = page.screenshot(type='png', full_page=False)
                
                if output_path:
                    with open(output_path, 'wb') as f:
                        f.write(screenshot_bytes)
                
                return screenshot_bytes
            finally:
                browser.close()
    
    async def render_to_image(self, slide_data: Dict[str, Any], output_path: Optional[str] = None) -> bytes:
        """
        渲染幻灯片为图片
        
        Args:
            slide_data: 幻灯片数据
            output_path: 输出路径（可选）
        
        Returns:
            图片字节数据
        """
        # 1. 渲染 HTML
        html_content = self.render_html(slide_data)
        
        # 2. HTML 转图片
        image_bytes = await self.html_to_image(html_content, output_path)
        
        return image_bytes


class LandPPTPPTXGenerator:
    """LandPPT PPTX 生成器"""
    
    def __init__(self, template_html: str):
        """
        初始化生成器
        
        Args:
            template_html: HTML 模板字符串
        """
        self.renderer = LandPPTTemplateRenderer(template_html)
        self.prs = Presentation()
        # 设置幻灯片尺寸为 16:9 (1280x720)
        self.prs.slide_width = Inches(10)
        self.prs.slide_height = Inches(5.625)
    
    async def add_slide_from_data(self, slide_data: Dict[str, Any], slide_number: int, total_slides: int):
        """
        从数据添加幻灯片
        
        Args:
            slide_data: 幻灯片数据，包含：
                - title: 标题
                - content: 内容（列表或 HTML）
                - layout: 布局类型
            slide_number: 当前页码
            total_slides: 总页数
        """
        # 准备渲染数据
        render_data = {
            "page_title": slide_data.get("title", ""),
            "main_heading": slide_data.get("title", ""),
            "page_content": self._format_content(slide_data.get("content", [])),
            "current_page_number": slide_number,
            "total_page_count": total_slides
        }
        
        # 渲染为图片
        image_bytes = await self.renderer.render_to_image(render_data)
        
        # 添加到 PPTX
        self._add_image_slide(image_bytes)
    
    def _format_content(self, content: Any) -> str:
        """
        格式化内容为 HTML
        
        Args:
            content: 内容（可以是列表、字符串或 HTML）
        
        Returns:
            HTML 字符串
        """
        if isinstance(content, list):
            # 列表转为 HTML 无序列表
            items = "".join([f"<li>{item}</li>" for item in content])
            return f'<ul class="content-points">{items}</ul>'
        elif isinstance(content, str):
            # 如果已经是 HTML，直接返回
            if content.strip().startswith('<'):
                return content
            # 否则包装在段落中
            return f'<p>{content}</p>'
        else:
            return str(content)
    
    def _add_image_slide(self, image_bytes: bytes):
        """
        添加图片幻灯片
        
        Args:
            image_bytes: 图片字节数据
        """
        # 使用空白布局
        blank_layout = self.prs.slide_layouts[6]
        slide = self.prs.slides.add_slide(blank_layout)
        
        # 添加图片，填满整个幻灯片
        image_stream = io.BytesIO(image_bytes)
        slide.shapes.add_picture(
            image_stream,
            left=0,
            top=0,
            width=self.prs.slide_width,
            height=self.prs.slide_height
        )
    
    async def generate_pptx(
        self,
        slides_data: List[Dict[str, Any]],
        output_path: str,
        title: str = "演示文稿"
    ) -> str:
        """
        生成 PPTX 文件
        
        Args:
            slides_data: 幻灯片数据列表
            output_path: 输出路径
            title: PPT 标题
        
        Returns:
            输出文件路径
        """
        total_slides = len(slides_data)
        
        logger.info(f"开始生成 PPTX: {title}, 共 {total_slides} 页")
        
        for i, slide_data in enumerate(slides_data, 1):
            logger.info(f"渲染第 {i}/{total_slides} 页: {slide_data.get('title', '')}")
            await self.add_slide_from_data(slide_data, i, total_slides)
        
        # 保存 PPTX
        self.prs.save(output_path)
        logger.info(f"PPTX 生成完成: {output_path}")
        
        return output_path


async def generate_pptx_with_landppt_template(
    template_html: str,
    slides_data: List[Dict[str, Any]],
    output_path: str,
    title: str = "演示文稿"
) -> str:
    """
    使用 LandPPT 模板生成 PPTX
    
    Args:
        template_html: HTML 模板字符串
        slides_data: 幻灯片数据列表，每个元素包含：
            - title: 标题
            - content: 内容（列表或 HTML 字符串）
            - layout: 布局类型（可选）
        output_path: 输出路径
        title: PPT 标题
    
    Returns:
        输出文件路径
    
    Example:
        slides_data = [
            {
                "title": "项目进展汇报",
                "content": ["完成核心功能", "用户增长200%", "营收突破500万"]
            },
            {
                "title": "核心成果",
                "content": "<div class='stats-grid'>...</div>"
            }
        ]
    """
    generator = LandPPTPPTXGenerator(template_html)
    return await generator.generate_pptx(slides_data, output_path, title)


# 便捷函数：从模板 ID 生成 PPTX
async def generate_pptx_from_template_id(
    template_id: str,
    slides_data: List[Dict[str, Any]],
    output_path: str,
    title: str = "演示文稿",
    db = None
) -> str:
    """
    从模板 ID 生成 PPTX
    
    Args:
        template_id: 模板 ID
        slides_data: 幻灯片数据列表
        output_path: 输出路径
        title: PPT 标题
        db: 数据库连接
    
    Returns:
        输出文件路径
    """
    if db is None:
        from database import get_database
        db = get_database()
    
    # 获取模板
    template = await db.ppt_templates.find_one({"template_id": template_id})
    if not template:
        # 使用默认模板
        template = await db.ppt_templates.find_one({"is_default": True})
    
    if not template:
        raise ValueError("未找到模板")
    
    template_html = template["html_template"]
    
    return await generate_pptx_with_landppt_template(
        template_html,
        slides_data,
        output_path,
        title
    )


if __name__ == "__main__":
    # 测试代码
    async def test():
        from database import get_database, connect_to_mongo
        
        await connect_to_mongo()
        db = get_database()
        
        # 测试数据
        slides_data = [
            {
                "title": "项目进展汇报",
                "content": ["完成核心功能开发", "用户增长200%", "营收突破500万"]
            },
            {
                "title": "核心成果",
                "content": ["技术架构升级", "性能提升3倍", "用户体验优化"]
            },
            {
                "title": "下一步计划",
                "content": ["扩展市场", "产品迭代", "团队建设"]
            }
        ]
        
        # 生成 PPTX
        output_path = "test_landppt.pptx"
        result = await generate_pptx_from_template_id(
            template_id="商务",
            slides_data=slides_data,
            output_path=output_path,
            title="项目进展汇报",
            db=db
        )
        
        print(f"生成成功: {result}")
    
    asyncio.run(test())
